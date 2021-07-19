import os

import xlwt
from docx import Document

path = input("请输入要搜索的路径(例如'D:\\Users\\shihaowang\\Projects\\Ten\\pyvrml\\demo\\search_word'):")

str_to_search = input("请输入要搜索的字段:")


def search_str_in_paragraph(str_to_search, paragraph):
    """
    本函数旨在特定段落中搜索指定字符串，并返回字符串索引位置。
    """
    str_index = paragraph.text.index(str_to_search)

    return str_index


def render_excel(title_list: [], content_rows: [[]], excel_name='excel.xls'):
    """
    生成EXCEL报表
    :param title_list: 标题
    :param content_rows: 内容
    :param excel_name: 文件名称
    """
    w = xlwt.Workbook(encoding='ascii')  # 创建一个工作簿
    ws = w.add_sheet('1')  # 创建一个工作表

    # 标题
    for j in range(0, len(title_list)):
        ws.write(0, j, title_list[j])

    # 行
    for i in range(0, len(content_rows)):
        content_row = content_rows[i]
        # 列
        for j in range(0, len(content_row)):
            ws.write(i + 1, j, content_row[j])

    w.save(excel_name)


def gen_excel(content_rows: []):
    title_list = ["text", "filename", "path"]
    render_excel(title_list=title_list,
                 content_rows=content_rows,
                 excel_name="word_keyword_search.xls")


if __name__ == '__main__':
    content_rows = []

    for root, dirs, files in os.walk(path, topdown=False):
        for file in files:
            filename = file.split(".")[0]
            ext = file.split(".")[1]
            if ext == "docx":
                doc = Document(os.path.join(root, file))
                for paragraph in doc.paragraphs:
                    if str_to_search in paragraph.text:
                        str_index = search_str_in_paragraph(str_to_search, paragraph)
                        text_value = paragraph.text[str_index - 10:str_index + 10]
                        path_value = os.path.join(root, file)
                        print(text_value, "-----", filename, "-----", path_value)
                        content_rows.append([text_value, filename, path_value])

    gen_excel(content_rows)
